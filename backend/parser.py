import fitz  # PyMuPDF
import docx
import os
import io
import olefile
import zipfile
import base64
import xml.etree.ElementTree as ET
from PIL import Image
from dotenv import load_dotenv

load_dotenv()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """PDF 파일에서 텍스트를 추출합니다."""
    text = ""
    try:
        pdf_document = fitz.open(stream=file_bytes, filetype="pdf")
        for page_num in range(len(pdf_document)):
            page = pdf_document.load_page(page_num)
            text += page.get_text() + "\n"
        return text.strip()
    except Exception as e:
        print(f"Error extracting PDF: {e}")
        return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """DOCX 파일에서 텍스트를 추출합니다."""
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        # 테이블 내 텍스트도 추출
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text += " | ".join(row_texts) + "\n"
        return text.strip()
    except Exception as e:
        print(f"Error extracting DOCX: {e}")
        return ""


def extract_text_from_hwp(file_bytes: bytes) -> str:
    """HWP 파일에서 텍스트를 추출합니다 (PrvText 기준)."""
    try:
        f = io.BytesIO(file_bytes)
        if not olefile.isOleFile(f):
            print("유효한 HWP 파일이 아닙니다.")
            return ""

        ole = olefile.OleFileIO(f)

        if ole.exists('PrvText'):
            stream = ole.openstream('PrvText')
            data = stream.read()
            # HWP PrvText는 utf-16le 형식입니다.
            text = data.decode('utf-16le', errors='ignore')
            return text.strip()
        else:
            print("본 HWP 파일은 텍스트 미리보기(PrvText)를 포함하지 않아 MVP 파서로 추출이 어렵습니다.")
            return ""
    except Exception as e:
        print(f"Error extracting HWP: {e}")
        return ""


def extract_text_from_hwpx(file_bytes: bytes) -> str:
    """HWPX(한글 표준 XML ZIP 포맷) 파일에서 텍스트를 추출합니다."""
    try:
        f = io.BytesIO(file_bytes)
        text_parts = []
        with zipfile.ZipFile(f, 'r') as z:
            # Contents/section0.xml, section1.xml 등 섹션 파일 탐색
            section_files = sorted([name for name in z.namelist() if name.startswith('Contents/section') and name.endswith('.xml')])
            for sname in section_files:
                xml_data = z.read(sname)
                root = ET.fromstring(xml_data)
                # XML 내 텍스트 노드 순회
                for elem in root.iter():
                    if elem.tag.endswith('}t') and elem.text:
                        text_parts.append(elem.text)
                    elif elem.tag.endswith('}p'):
                        text_parts.append("\n")
        return "".join(text_parts).strip()
    except Exception as e:
        print(f"Error extracting HWPX: {e}")
        return ""


def extract_text_from_image(file_bytes: bytes, filename: str) -> str:
    """이미지 파일(JPG, PNG, WEBP 등)에서 GPT-4o Vision OCR을 통해 고정밀 텍스트를 추출합니다."""
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key or api_key == "YOUR_API_KEY_HERE":
        print("[경고] OPENAI_API_KEY가 없어 이미지 텍스트를 추출할 수 없습니다.")
        return ""

    try:
        from openai import OpenAI

        # 이미지 리사이징 (해상도가 너무 크면 축소하여 전송 속도 및 비용 최적화)
        img = Image.open(io.BytesIO(file_bytes))
        max_dim = 2048
        if max(img.size) > max_dim:
            scale = max_dim / max(img.size)
            new_size = (int(img.size[0] * scale), int(img.size[1] * scale))
            img = img.resize(new_size, Image.Resampling.LANCZOS)

        # RGB 변환 (RGBA/Palette 모드 대응)
        if img.mode != 'RGB':
            img = img.convert('RGB')

        buf = io.BytesIO()
        img.save(buf, format='JPEG', quality=88)
        b64_img = base64.b64encode(buf.getvalue()).decode('utf-8')

        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "당신은 고정밀 계약서/법률 문서 전문 OCR AI입니다. "
                        "제공된 문서 이미지에 적힌 모든 문자(문서 제목, 조항 번호, 본문, 당사자 정보, "
                        "날짜, 금액, 서명란, 특약사항 등)를 빠짐없이 원문 그대로 정확하게 전사(transcribe)하십시오. "
                        "인사말이나 요약, 설명 등 사족은 일절 추가하지 말고 오직 추출된 문서 텍스트만 출력하십시오."
                    )
                },
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": "이 계약서/문서 이미지의 모든 내용을 누락 없이 텍스트로 정확히 추출해 주세요."
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/jpeg;base64,{b64_img}",
                                "detail": "high"
                            }
                        }
                    ]
                }
            ],
            temperature=0.0,
            max_tokens=4000
        )
        return (response.choices[0].message.content or "").strip()
    except Exception as e:
        print(f"Error extracting image text via Vision OCR: {e}")
        return ""


def parse_document(file_bytes: bytes, filename: str) -> str:
    """파일 확장자에 따라 적절한 파서를 호출하여 텍스트를 반환합니다.
    지원 확장자: .pdf, .docx, .hwp, .hwpx, .png, .jpg, .jpeg, .webp
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == '.pdf':
        return extract_text_from_pdf(file_bytes)
    elif ext == '.docx':
        return extract_text_from_docx(file_bytes)
    elif ext == '.hwp':
        return extract_text_from_hwp(file_bytes)
    elif ext == '.hwpx':
        return extract_text_from_hwpx(file_bytes)
    elif ext in ['.png', '.jpg', '.jpeg', '.webp']:
        return extract_text_from_image(file_bytes, filename)
    else:
        return ""


def parse_multiple_documents(file_items: list) -> str:
    """여러 파일(예: 여러 장의 계약서 사진, 분할 문서)을 병렬로 파싱하여 순서대로 합친 텍스트를 반환합니다.
    file_items: [(file_bytes, filename), ...]
    """
    if not file_items:
        return ""

    if len(file_items) == 1:
        return parse_document(file_items[0][0], file_items[0][1])

    import concurrent.futures

    # 여러 파일(특히 여러 장의 사진 OCR)을 병렬로 처리하여 속도 극대화
    max_workers = min(len(file_items), 8)
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        results = list(executor.map(lambda item: parse_document(item[0], item[1]), file_items))

    combined_parts = []
    for idx, (res_text, (_, fname)) in enumerate(zip(results, file_items)):
        if res_text.strip():
            combined_parts.append(f"=== [문서/페이지 {idx + 1}: {fname}] ===\n{res_text.strip()}")

    return "\n\n".join(combined_parts)

