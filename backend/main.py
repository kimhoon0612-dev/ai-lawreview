from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Request
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Any, Optional, List
from contextlib import asynccontextmanager
import uvicorn
import os
import io
import docx

from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from dotenv import load_dotenv

from parser import parse_document, parse_multiple_documents
from rag import analyze_contract_text, chat_with_law_db, summarize_document_text
from vector_store import (
    add_custom_documents, list_custom_documents, delete_custom_document,
    get_vector_store, search_documents
)
from law_fetcher import (
    sync_laws_from_api, get_law_api_key,
    search_ordinances, fetch_ordinance_articles, sync_ordinances_from_api
)
from auth import (
    register_user, login_user,
    get_current_user, get_optional_user,
)
from database import init_database, save_analysis, get_user_history, get_history_detail, delete_history

load_dotenv()

# 지원되는 파일 확장자 (PDF, Word, HWP/HWPX 및 사진/이미지 파일)
ALLOWED_EXTENSIONS = ('.pdf', '.docx', '.hwp', '.hwpx', '.png', '.jpg', '.jpeg', '.webp')

# 최대 파일 크기 (200MB로 대폭 확장)
MAX_FILE_SIZE = 200 * 1024 * 1024

# Rate Limiter 설정
limiter = Limiter(key_func=get_remote_address)

@asynccontextmanager
async def lifespan(app: FastAPI):
    # 서버 시작 시 DB 초기화
    init_database()
    port = int(os.getenv("PORT", 8001))
    print(f"[서버] AI LawReview API 서버 시작 (포트: {port})")
    yield
    print("[서버] AI LawReview API 서버 종료")

app = FastAPI(
    title="AI LawReview API",
    version="2.2.0",
    description="대한민국 법령 기반 AI 계약서 검토 및 다중 문서/사진 요약·검색 서비스",
    lifespan=lifespan,
)

# Rate Limiter 등록
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

# CORS 설정 (환경변수 및 Vercel 도메인 지원)
CORS_ORIGINS = os.getenv("CORS_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000,http://localhost:3001,http://127.0.0.1:3001").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in CORS_ORIGINS if origin.strip()],
    allow_origin_regex=r"https://.*\.vercel\.app",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# 데이터 모델
# ============================================================

class ReviewResponse(BaseModel):
    filename: str
    status: str
    message: str
    analysis: Optional[Any] = None
    extracted_text_preview: Optional[str] = None
    mock: bool = False
    history_id: Optional[int] = None

class ExportRequest(BaseModel):
    filename: str
    analysis: Any

class ChatRequest(BaseModel):
    message: str
    history: Optional[List[dict]] = None

class AuthRequest(BaseModel):
    email: str
    password: str
    name: Optional[str] = None

class SearchRequest(BaseModel):
    query: str
    scope: Optional[str] = "all"  # 'all' | 'law' | 'custom'
    limit: Optional[int] = 8


# ============================================================
# 헬스 체크
# ============================================================

@app.get("/")
def read_root():
    return {
        "message": "AI LawReview API is running.",
        "version": "2.2.0",
        "supported_formats": list(ALLOWED_EXTENSIONS),
        "max_upload_size_mb": 200,
        "multi_upload_supported": True,
        "port": int(os.getenv("PORT", 8001))
    }


# ============================================================
# 인증 API
# ============================================================

@app.post("/api/auth/register")
@limiter.limit("5/minute")
async def register(request: Request, body: AuthRequest):
    """회원가입"""
    if not body.name:
        raise HTTPException(status_code=400, detail="이름을 입력해 주세요.")
    result = register_user(body.email, body.password, body.name)
    if result["status"] == "error":
        raise HTTPException(status_code=400, detail=result["message"])
    return result


@app.post("/api/auth/login")
@limiter.limit("10/minute")
async def login(request: Request, body: AuthRequest):
    """로그인"""
    result = login_user(body.email, body.password)
    if result["status"] == "error":
        raise HTTPException(status_code=401, detail=result["message"])
    return result


@app.get("/api/auth/me")
async def get_me(user: dict = Depends(get_current_user)):
    """현재 로그인된 사용자 정보"""
    return {
        "status": "success",
        "user": {
            "id": user["id"],
            "email": user["email"],
            "name": user["name"],
        }
    }


# ============================================================
# 계약서 업로드 & AI 분석 (단일/다중 파일 및 다중 사진 지원)
# ============================================================

@app.post("/api/upload", response_model=ReviewResponse)
@limiter.limit("15/minute")
async def upload_document(
    request: Request,
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    sources: Optional[str] = Form(None),
    user: Optional[dict] = Depends(get_optional_user),
):
    all_files = []
    if files:
        all_files.extend(files)
    if file:
        all_files.append(file)

    if not all_files:
        raise HTTPException(status_code=400, detail="업로드할 파일이 없습니다.")

    # 사용자가 선택한 검토 기준 소스 파싱 (예: ["law", "ordinance", "custom"])
    target_sources = None
    if sources:
        try:
            import json
            if sources.strip().startswith("["):
                target_sources = json.loads(sources)
            else:
                target_sources = [s.strip() for s in sources.split(",") if s.strip()]
        except Exception:
            target_sources = ["law", "ordinance", "custom"]

    file_items = []
    total_size = 0
    display_name = ""

    for f in all_files:
        safe_name = os.path.basename(f.filename or "unknown")
        if not safe_name.lower().endswith(ALLOWED_EXTENSIONS):
            raise HTTPException(
                status_code=400,
                detail=f"지원하지 않는 파일 형식입니다: {safe_name} (지원: {', '.join(ALLOWED_EXTENSIONS)})"
            )
        fb = await f.read()
        total_size += len(fb)
        if len(fb) > MAX_FILE_SIZE or total_size > MAX_FILE_SIZE:
            raise HTTPException(status_code=413, detail="업로드 파일 총 크기가 200MB를 초과합니다.")
        file_items.append((fb, safe_name))

    first_name = file_items[0][1] if file_items else "unknown"
    if len(all_files) == 1:
        display_name = first_name
    else:
        display_name = f"{first_name} 외 {len(all_files) - 1}장 (총 {len(all_files)}페이지)"

    # 병렬 파싱 (다중 이미지 OCR 병렬 처리 포함)
    extracted_text = parse_multiple_documents(file_items)
    if not extracted_text.strip():
        raise HTTPException(status_code=400, detail="문서에서 텍스트를 추출할 수 없거나 빈 문서입니다.")

    # AI 검토 진행 (선택된 소스 기반 RAG 독소조항 분석)
    result = analyze_contract_text(extracted_text, target_sources=target_sources)

    if result.get("status") == "error":
         raise HTTPException(status_code=500, detail=result.get("message", "AI 분석 중 오류가 발생했습니다."))

    # 인증된 사용자인 경우 분석 이력 저장
    history_id = None
    if user:
        try:
            history_id = save_analysis(
                user_id=user["id"],
                filename=display_name,
                status=result["status"],
                analysis=result.get("analysis", {}),
                extracted_text_preview=result.get("extracted_text_preview", ""),
            )
        except Exception as e:
            print(f"[경고] 분석 이력 저장 실패: {e}")

    return {
        "filename": display_name,
        "status": result["status"],
        "message": result.get("message", "분석이 성공적으로 완료되었습니다."),
        "analysis": result.get("analysis"),
        "extracted_text_preview": result.get("extracted_text_preview"),
        "mock": result.get("mock", False),
        "history_id": history_id,
    }


# ============================================================
# 문서 자체 요약 API (Summary - 단일/다중 파일 및 다중 사진 지원)
# ============================================================

@app.post("/api/summarize")
@limiter.limit("20/minute")
async def summarize_endpoint(
    request: Request,
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    text: Optional[str] = Form(None),
):
    """문서(PDF, DOCX, HWP, HWPX, 다중 사진) 또는 직접 입력된 텍스트를 구조화 요약합니다."""
    extracted_text = ""
    doc_filename = "직접 입력 문서"

    all_files = []
    if files:
        all_files.extend(files)
    if file:
        all_files.append(file)

    if all_files:
        file_items = []
        total_size = 0
        for f in all_files:
            safe_name = os.path.basename(f.filename or "unknown")
            if not safe_name.lower().endswith(ALLOWED_EXTENSIONS):
                raise HTTPException(
                    status_code=400,
                    detail=f"지원하지 않는 파일 형식입니다: {safe_name} (지원: {', '.join(ALLOWED_EXTENSIONS)})"
                )
            fb = await f.read()
            total_size += len(fb)
            if len(fb) > MAX_FILE_SIZE or total_size > MAX_FILE_SIZE:
                raise HTTPException(status_code=413, detail="업로드 파일 총 크기가 200MB를 초과합니다.")
            file_items.append((fb, safe_name))

        first_name = file_items[0][1] if file_items else "unknown"
        if len(all_files) == 1:
            doc_filename = first_name
        else:
            doc_filename = f"{first_name} 외 {len(all_files) - 1}장 (총 {len(all_files)}페이지)"

        extracted_text = parse_multiple_documents(file_items)
    elif text and text.strip():
        extracted_text = text.strip()
    else:
        raise HTTPException(status_code=400, detail="분석할 파일이나 텍스트를 입력해 주세요.")

    if not extracted_text.strip():
        raise HTTPException(status_code=400, detail="문서에서 텍스트를 추출할 수 없거나 빈 문서입니다.")

    result = summarize_document_text(extracted_text)
    if result.get("status") == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "문서 요약 중 오류 발생"))

    result["filename"] = doc_filename
    return result


# ============================================================
# 법령 및 사내 자료 시맨틱 통합 검색 API (Search)
# ============================================================

@app.post("/api/search")
@limiter.limit("30/minute")
async def search_endpoint(request: Request, body: SearchRequest):
    """법령 DB 및 등록된 사내 참고자료를 시맨틱 통합 검색합니다."""
    if not body.query or not body.query.strip():
        raise HTTPException(status_code=400, detail="검색어를 입력해 주세요.")

    results = search_documents(
        query=body.query.strip(),
        scope=body.scope or "all",
        limit=body.limit or 8
    )

    return {
        "query": body.query.strip(),
        "scope": body.scope or "all",
        "total": len(results),
        "results": results
    }


@app.get("/api/search")
@limiter.limit("30/minute")
async def search_get_endpoint(request: Request, q: str, scope: str = "all", limit: int = 8):
    """GET 방식 시맨틱 검색 엔드포인트"""
    if not q or not q.strip():
        raise HTTPException(status_code=400, detail="검색어를 입력해 주세요.")

    results = search_documents(
        query=q.strip(),
        scope=scope,
        limit=limit
    )

    return {
        "query": q.strip(),
        "scope": scope,
        "total": len(results),
        "results": results
    }


# ============================================================
# 분석 이력 API
# ============================================================

@app.get("/api/history")
async def list_history(user: dict = Depends(get_current_user)):
    """사용자의 분석 이력 목록"""
    history = get_user_history(user["id"])
    return {"history": history, "total": len(history)}


@app.get("/api/history/{history_id}")
async def get_history_item(history_id: int, user: dict = Depends(get_current_user)):
    """특정 분석 이력 상세"""
    item = get_history_detail(history_id, user["id"])
    if not item:
        raise HTTPException(status_code=404, detail="이력을 찾을 수 없습니다.")
    return item


@app.delete("/api/history/{history_id}")
async def remove_history(history_id: int, user: dict = Depends(get_current_user)):
    """분석 이력 삭제"""
    success = delete_history(history_id, user["id"])
    if not success:
        raise HTTPException(status_code=404, detail="이력을 찾을 수 없거나 삭제할 수 없습니다.")
    return {"status": "success", "message": "이력이 삭제되었습니다."}


# ============================================================
# 리포트 내보내기 (Word .docx)
# ============================================================

@app.post("/api/export")
async def export_report(request: ExportRequest):
    try:
        doc = docx.Document()
        doc.add_heading('AI 법률 계약서 검토 리포트', 0)
        
        doc.add_paragraph(f'검토 대상 문서: {request.filename}')
        doc.add_heading('1. 전반적 조언', level=1)
        
        analysis = request.analysis or {}
        general_advice = analysis.get('general_advice', '제공된 일반 조언이 없습니다.')
        doc.add_paragraph(general_advice)
        
        doc.add_heading('2. 조항별 검토 상세', level=1)
        clauses = analysis.get('clauses', [])
        
        if not clauses:
            doc.add_paragraph("분석된 조항 내용이 없습니다.")
            
        for idx, clause in enumerate(clauses):
            risk = clause.get('risk_level', 'info')
            risk_text = "고위험" if risk == "high" else "주의" if risk == "warning" else "일반정보"
            
            p = doc.add_paragraph()
            p.add_run(f"[{risk_text}] {clause.get('clause_name', '')}").bold = True
            
            doc.add_paragraph("원본 조항:")
            doc.add_paragraph(clause.get('original_text', ''), style='List Bullet')
            
            doc.add_paragraph("AI 리뷰 내용:")
            doc.add_paragraph(clause.get('ai_review', ''), style='List Bullet')

            legal_basis = clause.get('legal_basis', '')
            if legal_basis:
                doc.add_paragraph(f"법적 근거: {legal_basis}")

            confidence = clause.get('confidence', '')
            if confidence:
                conf_text = "높음" if confidence == "high" else "보통" if confidence == "medium" else "낮음"
                doc.add_paragraph(f"분석 신뢰도: {conf_text}")
            
            doc.add_paragraph("-" * 40)

        # 면책 조항 추가
        doc.add_heading('면책 사항', level=1)
        doc.add_paragraph(
            '본 리포트는 AI가 제공하는 참고용 법률 정보이며, 공인된 법률 자문이 아닙니다. '
            '중요한 법률적 판단은 반드시 전문 변호사의 검토를 받으시기 바랍니다. '
            '본 서비스의 분석 결과에 의존하여 발생하는 어떠한 법적 불이익에 대해서도 책임을 지지 않습니다.'
        )
            
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        import urllib.parse
        encoded_filename = urllib.parse.quote(f"review_report_{request.filename}.docx")
        
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
        )
    except Exception as e:
        print(f"Export Error: {e}")
        raise HTTPException(status_code=500, detail="리포트 생성 중 오류가 발생했습니다.")


# ============================================================
# AI 법률 상담 채팅
# ============================================================

@app.post("/api/chat")
@limiter.limit("20/minute")
async def chat_endpoint(request: Request, body: ChatRequest):
    result = chat_with_law_db(body.message, body.history)
    if result.get("status") == "error":
        raise HTTPException(status_code=500, detail=result.get("reply", "채팅 중 오류가 발생했습니다."))
    return {"reply": result.get("reply")}


# ============================================================
# 사용자 정의 참고자료 관리 API (다중 파일 지원)
# ============================================================

@app.post("/api/custom-docs/upload")
async def upload_custom_document(
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    doc_name: str = Form(...),
    category: str = Form("other"),
):
    """사용자 정의 참고자료(회사 규정, 공공기관 규정 등)를 업로드합니다. 다중 사진 및 다중 파일 지원."""
    all_files = []
    if files:
        all_files.extend(files)
    if file:
        all_files.append(file)

    if not all_files:
        raise HTTPException(status_code=400, detail="업로드할 파일이 없습니다.")

    valid_categories = ["company_rule", "public_rule", "internal_guideline", "other"]
    if category not in valid_categories:
        raise HTTPException(status_code=400, detail=f"유효하지 않은 카테고리입니다. ({', '.join(valid_categories)})")

    file_items = []
    total_size = 0
    for f in all_files:
        safe_name = os.path.basename(f.filename or "unknown")
        if not safe_name.lower().endswith(ALLOWED_EXTENSIONS):
            raise HTTPException(
                status_code=400,
                detail=f"지원하지 않는 파일 형식입니다: {safe_name} (지원: {', '.join(ALLOWED_EXTENSIONS)})"
            )
        fb = await f.read()
        total_size += len(fb)
        if len(fb) > MAX_FILE_SIZE or total_size > MAX_FILE_SIZE:
            raise HTTPException(status_code=413, detail="업로드 파일 총 크기가 200MB를 초과합니다.")
        file_items.append((fb, safe_name))

    extracted_text = parse_multiple_documents(file_items)
    if not extracted_text.strip():
        raise HTTPException(status_code=400, detail="문서에서 텍스트를 추출할 수 없거나 빈 문서입니다.")

    result = add_custom_documents(extracted_text, doc_name, category)

    if result.get("status") == "error":
        raise HTTPException(status_code=500, detail=result.get("message", "참고자료 저장 중 오류가 발생했습니다."))

    return result


@app.get("/api/custom-docs")
async def get_custom_documents():
    """저장된 사용자 정의 참고자료 목록을 반환합니다."""
    docs = list_custom_documents()
    return {"documents": docs, "total": len(docs)}


@app.delete("/api/custom-docs/{doc_id}")
async def remove_custom_document(doc_id: str):
    """특정 사용자 정의 참고자료를 삭제합니다."""
    result = delete_custom_document(doc_id)

    if result.get("status") == "error":
        raise HTTPException(status_code=404, detail=result.get("message", "삭제 실패"))

    return result


# ============================================================
# 법령 DB 동기화 API
# ============================================================

@app.post("/api/law-sync")
async def sync_law_database():
    """국가법령정보센터 API를 통해 법령 DB를 동기화합니다."""
    try:
        vector_store = get_vector_store()
        if vector_store is None:
            raise HTTPException(status_code=500, detail="Vector Store 초기화 실패")

        result = sync_laws_from_api()
        documents = result.get("documents", [])
        if documents:
            doc_ids = [f"{doc.metadata.get('doc_id', 'law')}_{i}" for i, doc in enumerate(documents)]
            vector_store.add_documents(documents, ids=doc_ids)

        return {
            "status": result["status"],
            "message": result["message"],
            "count": len(documents),
            "has_api_key": get_law_api_key() is not None,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"동기화 중 오류: {str(e)}")


@app.get("/api/law-sync/status")
async def get_law_sync_status():
    """현재 법령 및 자치법규 DB 상태를 반환합니다."""
    try:
        vector_store = get_vector_store()
        if vector_store is None:
            return {"law_count": 0, "custom_count": 0, "ordinance_count": 0, "has_api_key": False}

        law_count = 0
        custom_count = 0
        ordinance_count = 0
        try:
            law_res = vector_store._collection.get(where={"source_type": "law"})
            law_count = len(law_res.get("ids", []))
        except Exception:
            pass
        try:
            custom_res = vector_store._collection.get(where={"source_type": "custom"})
            custom_count = len(custom_res.get("ids", []))
        except Exception:
            pass
        try:
            ordin_res = vector_store._collection.get(where={"source_type": "ordinance"})
            ordinance_count = len(ordin_res.get("ids", []))
        except Exception:
            pass

        return {
            "law_count": law_count,
            "custom_count": custom_count,
            "ordinance_count": ordinance_count,
            "has_api_key": get_law_api_key() is not None
        }
    except Exception:
        return {"law_count": 0, "custom_count": 0, "ordinance_count": 0, "has_api_key": False}


@app.post("/api/ordinances/sync")
async def sync_ordinances():
    """충청남도 천안시 및 천안시의회 자치법규를 수집하여 DB에 동기화합니다."""
    try:
        vector_store = get_vector_store()
        if vector_store is None:
            raise HTTPException(status_code=500, detail="Vector Store 초기화 실패")

        result = sync_ordinances_from_api()
        documents = result.get("documents", [])
        if documents:
            doc_ids = [f"{doc.metadata.get('doc_id', 'ordin')}_{i}" for i, doc in enumerate(documents)]
            vector_store.add_documents(documents, ids=doc_ids)

        return {
            "status": result["status"],
            "message": result["message"],
            "count": len(documents),
            "stats": result.get("stats", {})
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"자치법규 동기화 오류: {str(e)}")


@app.get("/api/ordinances/live-search")
async def live_search_ordinances(query: str = "천안시", display: int = 50):
    """국가법령정보센터에서 자치법규 및 의회 정보를 실시간 검색합니다."""
    api_key = get_law_api_key()
    if not api_key:
        # 키가 없을 경우 샘플 반환
        return {
            "query": query,
            "total": 3,
            "items": [
                {"name": "천안시 기업인 예우 및 기업활동 지원에 관한 조례", "mst": "sample1", "org": "충청남도 천안시", "type": "조례", "date": "20240101"},
                {"name": "천안시의회 회의 규칙", "mst": "sample2", "org": "천안시의회", "type": "규칙", "date": "20230621"},
                {"name": "천안시 소상공인 지원 조례", "mst": "sample3", "org": "충청남도 천안시", "type": "조례", "date": "20231215"}
            ]
        }

    results = search_ordinances(query, api_key, display=display)
    return {
        "query": query,
        "total": len(results),
        "items": results
    }


@app.get("/api/ordinances/articles")
async def get_ordinance_articles(mst: str):
    """특정 자치법규 일련번호(MST)의 전체 조문 내용을 가져옵니다."""
    api_key = get_law_api_key()
    if not api_key or mst.startswith("sample"):
        return {
            "mst": mst,
            "articles": [
                {"law_name": "천안시 조례 샘플", "org": "충청남도 천안시", "article": "제1조(목적)", "content": "이 조례는 천안시 관내 공공 행정 및 지원을 규정함을 목적으로 한다."}
            ]
        }

    articles = fetch_ordinance_articles(mst, api_key)
    return {
        "mst": mst,
        "total": len(articles),
        "articles": articles
    }


@app.get("/api/laws/catalog")
async def get_laws_catalog_endpoint():
    """DB에 적재된 모든 국가 법령 및 천안시 자치법규의 전체 목록과 조항 수를 반환합니다."""
    import sqlite3
    db_path = os.path.join(os.path.dirname(__file__), "chroma_db", "chroma.sqlite3")
    if not os.path.exists(db_path):
        return {"national": [], "ordinance": [], "total_laws": 0, "total_articles": 0}
    try:
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()
        sql = """
            SELECT 
                m_src.string_value AS source_name,
                COALESCE(m_type.string_value, 'law') AS source_type,
                COALESCE(m_org.string_value, '') AS org,
                COUNT(DISTINCT m_src.id) AS article_count
            FROM embedding_metadata m_src
            LEFT JOIN embedding_metadata m_type ON m_src.id = m_type.id AND m_type.key = 'source_type'
            LEFT JOIN embedding_metadata m_org ON m_src.id = m_org.id AND m_org.key = 'org'
            WHERE m_src.key = 'source'
            GROUP BY m_src.string_value
            ORDER BY article_count DESC
        """
        cur.execute(sql)
        national = []
        ordinances = []
        for name, stype, org, count in cur.fetchall():
            item = {
                "name": name,
                "type": stype,
                "org": org or ("충청남도 천안시" if stype == "ordinance" else "대한민국"),
                "count": count
            }
            if stype == "ordinance":
                ordinances.append(item)
            else:
                national.append(item)
        conn.close()
        return {
            "national": national,
            "ordinance": ordinances,
            "total_laws": len(national) + len(ordinances),
            "total_articles": sum(x["count"] for x in national) + sum(x["count"] for x in ordinances)
        }
    except Exception as e:
        print(f"Error fetching laws catalog: {e}")
        return {"national": [], "ordinance": [], "total_laws": 0, "total_articles": 0}


@app.get("/api/laws/full-text")
async def get_law_full_text_endpoint(source_name: str):
    """특정 법률 또는 자치법규의 제1조부터 마지막 조항까지 전체 전문을 반환합니다."""
    import sqlite3
    db_path = os.path.join(os.path.dirname(__file__), "chroma_db", "chroma.sqlite3")
    if not os.path.exists(db_path):
        return {"source_name": source_name, "total": 0, "articles": []}
    try:
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()
        sql = """
            SELECT 
                COALESCE(m_art.string_value, '') AS article,
                COALESCE(m_doc.string_value, '') AS content,
                COALESCE(m_org.string_value, '') AS org
            FROM embedding_metadata m_src
            LEFT JOIN embedding_metadata m_art ON m_src.id = m_art.id AND m_art.key = 'article'
            LEFT JOIN embedding_metadata m_doc ON m_src.id = m_doc.id AND m_doc.key = 'chroma:document'
            LEFT JOIN embedding_metadata m_org ON m_src.id = m_org.id AND m_org.key = 'org'
            WHERE m_src.key = 'source' AND (m_src.string_value = ? OR m_src.string_value LIKE ?)
            ORDER BY m_src.id ASC
        """
        cur.execute(sql, (source_name.strip(), f"%{source_name.strip()}%"))
        articles = []
        org_name = ""
        for art, content, org in cur.fetchall():
            if org and not org_name:
                org_name = org
            articles.append({
                "article": art,
                "content": content
            })
        conn.close()
        return {
            "source_name": source_name,
            "org": org_name or "대한민국",
            "total": len(articles),
            "articles": articles
        }
    except Exception as e:
        print(f"Error fetching law full text for {source_name}: {e}")
        return {"source_name": source_name, "total": 0, "articles": []}


if __name__ == "__main__":
    port = int(os.getenv("PORT", 8001))
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=True)
